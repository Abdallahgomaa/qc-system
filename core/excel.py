"""تصدير جدول إلى ملف Excel منسّق (RTL، رأس كحلي) كما في المتطلبات (FR-7)."""
from __future__ import annotations

import io

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NAVY = "0F2540"
HEAD_FONT = Font(color="FFFFFF", bold=True)
HEAD_FILL = PatternFill("solid", fgColor=NAVY)


def _text(value) -> str:
    """تمثيل نصي آمن لأي خلية (سلسلة فارغة للقيم الناقصة/NaN).

    لا نعتمد على ``astype(str)`` لأن سلوكه غير موثوق مع أعمدة NaN في بعض إصدارات
    pandas (تعيد float داخل السلسلة).
    """
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value)


def df_to_excel(df: pd.DataFrame, sheet_name: str = "CGES") -> bytes:
    """يعيد محتوى ملف ``.xlsx`` جاهزاً للتنزيل.

    ينسّق الورقة من اليمين لليسار، ويجعل رأس العمود بخلفية كحلية وخط أبيض،
    ويضبط عرض الأعمدة بحسب أطول قيمة فيها.
    """
    try:
        sheet = (sheet_name or "CGES")[:31] or "CGES"
        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name=sheet)
            ws = writer.sheets[sheet]
            ws.sheet_view.rightToLeft = True

            # نصل للأعمدة موضعياً لتفادي مشاكل تكرار الأسماء، ونحسب العرض بتحويل آمن للنص.
            for idx in range(1, len(df.columns) + 1):
                header = ws.cell(row=1, column=idx)
                header.fill = HEAD_FILL
                header.font = HEAD_FONT
                header.alignment = Alignment(horizontal="center", vertical="center")

                longest = max((len(_text(v)) for v in df.iloc[:, idx - 1].tolist()), default=0)
                ws.column_dimensions[get_column_letter(idx)].width = max(12, min(42, longest + 4))
        return buf.getvalue()
    except Exception as e:  # noqa: BLE001
        raise RuntimeError(f"تعذّر إنشاء ملف Excel: {e}") from e
